# -*- coding: utf-8 -*-
"""
Generate the Official NexusKV System Design Presentation (Minimalist, Clean, Sleek)
Spacious typography, uncluttered cards, elegant dark slate palette.
Architect: Ali Hasanifard
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def set_shape_bg(shape, hex_color):
    fill = shape.fill
    fill.solid()
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    fill.fore_color.rgb = RGBColor(r, g, b)

def apply_rtl(p):
    p_pr = p._p.get_or_add_pPr()
    p_pr.set('rtl', '1')

def add_header(slide, title_text, category_text="NEXUS-KV // ARCHITECTURE"):
    header_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_cat = tf.paragraphs[0]
    p_cat.alignment = PP_ALIGN.RIGHT
    apply_rtl(p_cat)
    r_cat = p_cat.add_run()
    r_cat.text = category_text
    r_cat.font.name = "JetBrains Mono"
    r_cat.font.size = Pt(11)
    r_cat.font.bold = True
    r_cat.font.color.rgb = RGBColor(99, 102, 241) # Indigo Accent
    
    p_title = tf.add_paragraph()
    p_title.space_before = Pt(4)
    p_title.alignment = PP_ALIGN.RIGHT
    apply_rtl(p_title)
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.name = "Vazirmatn"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(248, 250, 252)

def add_footer(slide, current_idx, total_slides):
    footer_box = slide.shapes.add_textbox(Inches(0.9), Inches(6.8), Inches(11.5), Inches(0.4))
    tf = footer_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    apply_rtl(p)
    r = p.add_run()
    r.text = f"NexusKV Distributed Storage Engine  •  Ali Hasanifard  •  {current_idx:02d} / {total_slides:02d}"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(100, 116, 139)

def add_card(slide, left, top, width, height, title, items, tag=None):
    card = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    set_shape_bg(card, "11131a")
    card.line.color.rgb = RGBColor(38, 42, 58)
    card.line.width = Pt(1.0)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.28)
    tf.margin_bottom = Inches(0.25)
    
    p_title = tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.RIGHT
    apply_rtl(p_title)
    
    if tag:
        r_b = p_title.add_run()
        r_b.text = f"[{tag}] "
        r_b.font.name = "JetBrains Mono"
        r_b.font.size = Pt(11)
        r_b.font.color.rgb = RGBColor(99, 102, 241)
        
    r_t = p_title.add_run()
    r_t.text = title
    r_t.font.name = "Vazirmatn"
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(241, 245, 249)
    
    for item in items:
        p_item = tf.add_paragraph()
        p_item.space_before = Pt(10)
        p_item.alignment = PP_ALIGN.RIGHT
        apply_rtl(p_item)
        r_bullet = p_item.add_run()
        r_bullet.text = "— "
        r_bullet.font.name = "JetBrains Mono"
        r_bullet.font.size = Pt(11)
        r_bullet.font.color.rgb = RGBColor(99, 102, 241)
        
        r_text = p_item.add_run()
        r_text.text = item
        r_text.font.name = "Vazirmatn"
        r_text.font.size = Pt(12)
        r_text.font.color.rgb = RGBColor(203, 213, 225)

def add_code_card(slide, left, top, width, height, title, code_snippet):
    card = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    set_shape_bg(card, "090a0f")
    card.line.color.rgb = RGBColor(38, 42, 58)
    card.line.width = Pt(1.0)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    
    p_title = tf.paragraphs[0]
    p_title.alignment = PP_ALIGN.RIGHT
    apply_rtl(p_title)
    r_t = p_title.add_run()
    r_t.text = title
    r_t.font.name = "JetBrains Mono"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(99, 102, 241)
    
    p_code = tf.add_paragraph()
    p_code.space_before = Pt(10)
    p_code.alignment = PP_ALIGN.LEFT
    r_c = p_code.add_run()
    r_c.text = code_snippet
    r_c.font.name = "JetBrains Mono"
    r_c.font.size = Pt(11)
    r_c.font.color.rgb = RGBColor(165, 180, 252)

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    TOTAL_SLIDES = 10
    
    def make_base_slide(chapter_tag, title_text, idx):
        s = prs.slides.add_slide(blank_layout)
        bg = s.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        set_shape_bg(bg, "090a0f")
        bg.line.fill.background()
        add_header(s, title_text, chapter_tag)
        add_footer(s, idx, TOTAL_SLIDES)
        return s

    # SLIDE 1: Cover (Clean, Spacious Minimal)
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    set_shape_bg(bg1, "090a0f")
    bg1.line.fill.background()
    
    box1 = s1.shapes.add_shape(1, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    set_shape_bg(box1, "11131a")
    box1.line.color.rgb = RGBColor(38, 42, 58)
    box1.line.width = Pt(1.0)
    
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    apply_rtl(p)
    r = p.add_run()
    r.text = "NexusKV"
    r.font.name = "JetBrains Mono"
    r.font.size = Pt(36)
    r.font.bold = True
    r.font.color.rgb = RGBColor(248, 250, 252)
    
    p2 = tf1.add_paragraph()
    p2.space_before = Pt(8)
    p2.alignment = PP_ALIGN.CENTER
    apply_rtl(p2)
    r2 = p2.add_run()
    r2.text = "Distributed Raft Consensus & LSM-Tree Engine in Pure Rust"
    r2.font.name = "JetBrains Mono"
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(99, 102, 241)
    
    p3 = tf1.add_paragraph()
    p3.space_before = Pt(24)
    p3.alignment = PP_ALIGN.CENTER
    apply_rtl(p3)
    r3 = p3.add_run()
    r3.text = "سیستم پایگاه داده توزیع‌شده با توان پردازش +۱۰۲ هزار درخواست بر ثانیه و سازگاری کامل با پروتکل Redis RESP"
    r3.font.name = "Vazirmatn"
    r3.font.size = Pt(14)
    r3.font.color.rgb = RGBColor(203, 213, 225)
    
    p4 = tf1.add_paragraph()
    p4.space_before = Pt(40)
    p4.alignment = PP_ALIGN.CENTER
    apply_rtl(p4)
    r4 = p4.add_run()
    r4.text = "طراح و معمار سیستم: علی حسنی‌فرد  •  github.com/alihasanifard-web/nexus-kv"
    r4.font.name = "Vazirmatn"
    r4.font.size = Pt(12)
    r4.font.color.rgb = RGBColor(148, 163, 184)
    add_footer(s1, 1, TOTAL_SLIDES)

    # SLIDE 2: High-Level Architecture
    s2 = make_base_slide("SECTION 01  //  ARCHITECTURE", "معماری سه‌لایه سیستم پایگاه داده توزیع‌شده NexusKV", 2)
    add_card(s2, 0.9, 1.8, 5.6, 4.6, "لایه‌های سه‌گانه سیستم", [
        "لایه ۱ (شبکه): سرور ناهمگام Tokio با پشتیبانی از پروتکل استاندارد Redis RESP v2.",
        "لایه ۲ (اجماع): ماشین حالت Raft برای تکثیر لاگ و مدیریت لیدر و فالوورها.",
        "لایه ۳ (ذخیره‌سازی): ساختار LSM-Tree با MemTable، WAL و SSTable پایدار.",
        "تراکنش‌های ۲ فازی 2PC با ساعت منطقی هیبریدی HLC."
    ], "CORE")
    add_code_card(s2, 6.8, 1.8, 5.6, 4.6, "EXECUTION WRITE PATH",
        "[Client: SET key val]\n         ↓ (RESP Protocol)\n[Tokio Async TCP Worker]\n         ↓\n[Raft Propose -> Append Log]\n         ↓ (Quorum 2/3 Nodes)\n[Commit Log Entry]\n         ↓\n[WAL Append (CRC32)] -> [MemTable Put]\n         ↓ (Size >= Threshold)\n[Flush to Disk: SSTable]"
    )

    # SLIDE 3: LSM-Tree Storage
    s3 = make_base_slide("SECTION 02  //  LSM-TREE", "موتور ذخیره‌سازی درخت LSM در برابر B-Tree", 3)
    add_card(s3, 0.9, 1.8, 5.6, 4.6, "مزایای LSM-Tree در نوشتن سنگین", [
        "تبدیل نوشتن‌های تصادفی دیسک به نوشتن ترتیبی متوالی (Sequential I/O).",
        "حذف گلوگاه قفل با استفاده از بافر BTreeMap در حافظه رم.",
        "حذف رکوردها با نشانه‌گذاری Tombstone بدون کاهش سرعت نوشتن.",
        "فرآیند یکپارچه‌سازی و فشرده‌سازی در پس‌زمینه (Background Compaction)."
    ], "STORAGE")
    add_card(s3, 6.8, 1.8, 5.6, 4.6, "اجزای داخلی موتور ذخیره‌سازی", [
        "MemTable: بافر حافظه رم مرتب‌شده با پایش لحظه‌ای حجم بایت‌ها.",
        "WAL: لاگ مکتوب دیسک با فریم‌های اعتبارسنجی شده توسط CRC32.",
        "SSTable: فایل‌های تغییرناپذیر دیسک همراه با ایندکس تنک.",
        "فیلتر بلوم: ساختار احتمالی دابل-هشینگ با ضریب خطای ۱٪ جهت کاهش رجوع به دیسک."
    ], "MODULES")

    # SLIDE 4: Raft Consensus
    s4 = make_base_slide("SECTION 03  //  CONSENSUS", "الگوریتم اجماع Raft و مقاومت در برابر قطعی شبکه", 4)
    add_card(s4, 0.9, 1.8, 5.6, 4.6, "ماشین حالت و قوانین اجماع", [
        "نقش‌های سیستم: لیدر (Leader)، پیرو (Follower) و نامزد (Candidate).",
        "تایمر تصادفی انتخابات (۱۵۰ تا ۳۰۰ میلی‌ثانیه) جهت جلوگیری از تساوی آرا.",
        "ضربان قلب دوره‌ای (Heartbeat) هر ۵۰ میلی‌ثانیه توسط لیدر برای حفظ رهبری.",
        "حد نصاب اکثریت (Quorum = N/2 + 1) برای تایید نهایی داده‌ها."
    ], "RAFT")
    add_code_card(s4, 6.8, 1.8, 5.6, 4.6, "RPC PROTOCOL SPECIFICATION",
        "// RPC درخواست رای\nRequestVoteArgs {\n    term: u64,\n    candidate_id: NodeId,\n    last_log_index: u64,\n    last_log_term: u64\n}\n\n// RPC تکثیر لاگ و هارت‌بیت\nAppendEntriesArgs {\n    term: u64,\n    leader_id: NodeId,\n    entries: Vec<LogEntry>,\n    leader_commit: u64\n}"
    )

    # SLIDE 5: Redis Compatibility
    s5 = make_base_slide("SECTION 04  //  PROTOCOL", "سازگاری کامل با پروتکل استاندارد Redis RESP", 5)
    add_card(s5, 0.9, 1.8, 5.6, 4.6, "کلاینت‌های بدون درایور اختصاصی", [
        "اتصال بدون تغییر کدهای موجود با redis-cli، redis-py، ioredis و redis-go.",
        "پارس استریم بدون کپی حافظه با بافرهای BytesMut.",
        "دستورات استاندارد: PING، SET، GET، DEL، ROLE و FLUSHALL.",
        "اسکن بازه‌ای داده‌ها با دستورات پیشرفته KEYS pattern و SCAN cursor."
    ], "RESP WIRE")
    add_code_card(s5, 6.8, 1.8, 5.6, 4.6, "OFFICIAL PYTHON CLIENT",
        "from nexus_kv import NexusClient\n\nwith NexusClient(port=6379) as client:\n    client.set('user:101', 'Ali Hasanifard')\n    val = client.get('user:101')\n    print('Value:', val)\n\n    # جستجوی کلیدها\n    keys = client.keys('user:*')\n    print('Matched:', keys)"
    )

    # SLIDE 6: Multi-Raft & Sharding
    s6 = make_base_slide("SECTION 05  //  SCALE", "معماری چندگروهه Multi-Raft و تقسیم خودکار Region", 6)
    add_card(s6, 0.9, 1.8, 5.6, 4.6, "افراز افقی داده‌ها (Dynamic Sharding)", [
        "تقسیم فضای کلیدها به بازه‌های مشخص به نام Region.",
        "تخصیص گروه مستقل Raft برای هر Region جهت توزیع بار نوشتن.",
        "مقیاس‌پذیری خطی سیستم تا حجم ترابایتی بدون افت کارایی.",
        "عملیات split_region برای دو نیم کردن خودکار ناحیه‌های پرحجم."
    ], "MULTI-RAFT")
    add_code_card(s6, 6.8, 1.8, 5.6, 4.6, "REGION ROUTING ENGINE",
        "let router = ShardRouter::new();\n\n// تقسیم خودکار ریجن ۱ در کلید 'm'\nrouter.split_region(1, b'm'.to_vec(), 2)?;\n\n// هدایت خودکار کلیدها به ریجن مربوطه\nlet r1 = router.locate_region(b'apple'); // ریجن ۱\nlet r2 = router.locate_region(b'zebra'); // ریجن ۲"
    )

    # SLIDE 7: Distributed ACID 2PC
    s7 = make_base_slide("SECTION 06  //  TRANSACTIONS", "پروتکل تراکنش دو فازی ۲PC و ساعت منطقی هیبریدی", 7)
    add_card(s7, 0.9, 1.8, 5.6, 4.6, "تضمین کامل اصول ACID", [
        "ساعت منطقی هیبریدی (HLC) برای ترتیب‌دهی دقیق زمانی تراکنش‌ها.",
        "کنترل همزمانی خوش‌بینانه (OCC) جهت جلوگیری از تضادهای همزمان.",
        "فاز ۱ (Prewrite): اخذ قفل‌های محلی و ثبت تغییرات موقت.",
        "فاز ۲ (Commit): تخصیص زمان نهایی و آزادسازی اتمیک قفل‌ها."
    ], "2PC")
    add_code_card(s7, 6.8, 1.8, 5.6, 4.6, "ATOMIC TRANSACTION SPEC",
        "let mut txn = Transaction::begin();\n\n// انتقال اتمیک موجودی حساب‌ها\ntxn.set(b'acc:alice'.to_vec(), b'900'.to_vec());\ntxn.set(b'acc:bob'.to_vec(), b'1100'.to_vec());\n\n// فاز اول: قفل‌گذاری\ncoordinator.prewrite(&txn)?;\n\n// فاز دوم: تایید و اعمال نهایی\nlet commit_ts = coordinator.commit(&mut txn)?;"
    )

    # SLIDE 8: Benchmarks & Latency
    s8 = make_base_slide("SECTION 07  //  BENCHMARKS", "نتایج بنچمارک رسمی و تحلیل تاخیر صدک‌ها", 8)
    add_card(s8, 0.9, 1.8, 5.6, 4.6, "تحلیل کارایی و نرخ پردازش", [
        "سنجش عملیاتی با ابزار چندترده nexus-bench روی حافظه NVMe.",
        "بیش از ۱۰۲ هزار درخواست در ثانیه تحت ۵۰ ورکر همزمان.",
        "میانگین تاخیر میانه (P50) زیر ۳۴۰ میکروثانیه.",
        "حداکثر تاخیر صدک ۹۹ (P99) در مرز ۱.۲ میلی‌ثانیه."
    ], "PERFORMANCE")
    add_card(s8, 6.8, 1.8, 5.6, 4.6, "جدول تاخیرات رسمی", [
        "• P50  Latency (Median):    340 µs",
        "• P90  Latency:             780 µs",
        "• P95  Latency:             960 µs",
        "• P99  Latency:           1,240 µs (1.2 ms)",
        "• P99.9 Latency:          2,100 µs (2.1 ms)",
        "• نرخ توان عملیاتی:        102,100 QPS"
    ], "METRICS")

    # SLIDE 9: Chaos Testing & K8s
    s9 = make_base_slide("SECTION 08  //  RESILIENCE", "تست آشوب خودکار (Chaos Test) و استقرار ابری", 9)
    add_card(s9, 0.9, 1.8, 5.6, 4.6, "نتایج تست آشوب (Chaos Testing)", [
        "اسکریپت اختصاصی scripts/chaos_test.py.",
        "شلیک سیگنال SIGKILL به نود لیدر در حین ورود داده‌ها.",
        "بازیابی سرد بدون افت حتی ۱ بایت از لاگ‌های مکتوب WAL.",
        "تداوم سرویس‌دهی نودها پس از انتخاب لیدر جدید."
    ], "CHAOS LAB")
    add_card(s9, 6.8, 1.8, 5.6, 4.6, "استقرار در محیط ابری", [
        "ایمیج Docker سبک کمتر از ۲۰ مگابایت با بیلد Alpine.",
        "مانیفست رسمی Kubernetes StatefulSet با ۳ رپلیکا.",
        "سرویس Headless جهت برقراری شبکه ارتباطی داخلی Raft.",
        "اندپوینت متریک استاندارد پرومتئوس با دستور METRICS."
    ], "KUBERNETES")

    # SLIDE 10: Conclusion & Author
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    set_shape_bg(bg10, "090a0f")
    bg10.line.fill.background()
    add_header(s10, "جمع‌بندی و هویت پروژه NexusKV", "CREDITS & SUMMARY")
    add_footer(s10, 10, TOTAL_SLIDES)
    
    box_c = s10.shapes.add_shape(1, Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.5))
    set_shape_bg(box_c, "11131a")
    box_c.line.color.rgb = RGBColor(38, 42, 58)
    box_c.line.width = Pt(1.0)
    
    tfc = box_c.text_frame
    tfc.word_wrap = True
    
    p1 = tfc.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    apply_rtl(p1)
    r1 = p1.add_run()
    r1.text = "NexusKV: Distributed High-Performance Storage"
    r1.font.name = "JetBrains Mono"
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(248, 250, 252)
    
    p2 = tfc.add_paragraph()
    p2.space_before = Pt(16)
    p2.alignment = PP_ALIGN.CENTER
    apply_rtl(p2)
    r2 = p2.add_run()
    r2.text = "کلاینت تعاملی nexus-cli • درایور رسمی پایتون • خط لوله CI گیت‌هاب • وایت‌پیپر و مستندات کامل"
    r2.font.name = "Vazirmatn"
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(203, 213, 225)
    
    p3 = tfc.add_paragraph()
    p3.space_before = Pt(36)
    p3.alignment = PP_ALIGN.CENTER
    apply_rtl(p3)
    r3 = p3.add_run()
    r3.text = "طراح و معمار سیستم:"
    r3.font.name = "Vazirmatn"
    r3.font.size = Pt(15)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(99, 102, 241)
    
    p4 = tfc.add_paragraph()
    p4.space_before = Pt(10)
    p4.alignment = PP_ALIGN.CENTER
    apply_rtl(p4)
    r4 = p4.add_run()
    r4.text = "علی حسنی‌فرد (Ali Hasanifard)\nEmail: hasanifardalipa@gmail.com  •  GitHub: github.com/alihasanifard-web/nexus-kv"
    r4.font.name = "JetBrains Mono"
    r4.font.size = Pt(13)
    r4.font.color.rgb = RGBColor(148, 163, 184)

    out_file = r"C:\Users\V I S I O N\Desktop\معرفی_معماری_NexusKV.pptx"
    prs.save(out_file)
    print(f"[OK] Minimalist Presentation saved to: {out_file}")

if __name__ == "__main__":
    create_deck()
